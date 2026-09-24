"""内容解析层：本地文件、网页、视频字幕、全网检索。

所有解析结果统一输出为 Markdown 文本，便于直接落入知识库。
视频转图文 / 图片 OCR 依赖外部二进制（ffmpeg / whisper / tesseract），
放在 Phase 8 补齐，此处仅给出可识别的占位逻辑。
"""
import logging
import os
import tempfile
from pathlib import Path

import requests
from bs4 import BeautifulSoup

# 网页解析实现已独立成 core/article_parser.py（保真优先：编码判定 / 正文容器优先 /
# 结构完整性校验 / 全结构 HTML→Markdown / 取回护栏）。这里只保留对外入口。
from core.article_parser import (HEADERS, WebParseError, parse_article,
                                 resolve_encoding)

logger = logging.getLogger('kb')

IMAGE_EXTS = {'.jpg', '.jpeg', '.png', '.bmp', '.gif', '.webp'}


# ---------------------------------------------------------------------------
# 本地文件解析
# ---------------------------------------------------------------------------
def parse_local_file(file_obj, orig_name, upload_url=''):
    """解析上传的本地文件 -> (title, markdown, meta_dict)。

    file_obj: Django UploadedFile（具备 .read() / .name）。
    """
    ext = Path(orig_name).suffix.lower()
    name = Path(orig_name).stem or '未命名'

    # 防御性复位：调用方可能已消费过流（如先 f.chunks() 落盘），
    # 不复位会导致 .md/.txt/.docx/.pdf/.pptx 等分支读到 EOF 而内容为空。
    try:
        file_obj.seek(0)
    except Exception:
        pass

    if ext in ('.md', '.markdown', '.txt'):
        raw = file_obj.read()
        text = raw.decode('utf-8', errors='ignore')
        return name, text, {'ext': ext}

    if ext == '.docx':
        from docx import Document
        doc = Document(file_obj)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        md = '\n\n'.join(lines)
        return name, md, {'ext': ext}

    if ext == '.pdf':
        import PyPDF2
        reader = PyPDF2.PdfReader(file_obj)
        pages = [ (p.extract_text() or '') for p in reader.pages ]
        md = '\n\n'.join(pages)
        return name, md, {'ext': ext, 'pages': len(reader.pages)}

    if ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(file_obj)
        blocks = []
        for i, slide in enumerate(prs.slides, 1):
            blocks.append(f'\n## 第 {i} 页')
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    blocks.append(shape.text_frame.text.strip())
        return name, '\n'.join(blocks), {'ext': ext, 'slides': len(prs.slides)}

    if ext == '.doc':
        raise ValueError('旧版 .doc 需要 LibreOffice 转换，当前版本暂不支持，请另存为 .docx。')

    if ext in IMAGE_EXTS:
        # 落盘临时文件后调用 OCR（Tesseract），缺失则给出提示
        tmp = tempfile.NamedTemporaryFile(suffix=ext, delete=False)
        file_obj.seek(0)
        tmp.write(file_obj.read())
        tmp.close()
        from core import media
        ocr_md = ''
        try:
            _, ocr_md = media.ocr_image_to_markdown(tmp.name, name)
        except Exception as e:
            logger.error('ocr error: %s', e)
        try:
            os.unlink(tmp.name)
        except Exception:
            pass
        md = (f'![{name}]({upload_url or orig_name})\n\n'
              f'> 图片 OCR 识别结果（已结构化清洗：修正误识、统一格式、过滤无效内容）：\n\n'
              f'{ocr_md or "（OCR 暂不可用，请安装 Tesseract-OCR 并加入 PATH）"}')
        return name, md, {'ext': ext, 'ocr': 'done' if ocr_md else 'unavailable'}

    raise ValueError(f'暂不支持的文件类型：{ext}')


# ---------------------------------------------------------------------------
# 网页解析（普通网页 / 公众号 / B站图文 / 技术文档）
# ---------------------------------------------------------------------------
def parse_web_page(url, platform='auto'):
    """抓取网页正文，返回 ``(title, markdown)``。

    失败时抛 ``WebParseError``（``ValueError`` 子类，带 ``reason`` 机器可读原因），
    调用方据此把链接归档到链接库并由 ``_fail_reason_from_exc`` 归类。

    真正的实现在 ``core.article_parser``：公众号走 ``#js_content``、B 站图文走
    开放接口、其余走正文容器（Readability 仅兜底且需通过结构完整性校验），
    并把代码块 / 图片 / 表格 / 列表完整转成 Markdown。
    """
    result = parse_article(url, platform)
    return result['title'], result['markdown']


# ---------------------------------------------------------------------------
# 全网检索（国内可访问的 Bing 结果解析，失败返回空列表）
# ---------------------------------------------------------------------------
def search_web(query, limit=10):
    try:
        url = 'https://www.bing.com/search?q=' + requests.utils.quote(query)
        r = requests.get(url, headers=HEADERS, timeout=15)
        # 不走 r.text：Bing 偶尔不带 charset，requests 会退回 ISO-8859-1 导致乱码
        r.encoding = resolve_encoding(r.content, r.headers.get('Content-Type') or '')
        soup = BeautifulSoup(r.text, 'html.parser')
        results = []
        for li in soup.select('li.b_algo')[:limit]:
            h = li.find('h2')
            if not h:
                continue
            a = h.find('a')
            results.append({
                'title': h.get_text(strip=True),
                'url': a.get('href') if a else '',
                'snippet': (li.find('p').get_text(strip=True) if li.find('p') else ''),
            })
        return results
    except Exception as e:
        logger.error('search_web failed: %s', e)
        return []
