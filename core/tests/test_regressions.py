import io
import json
import tempfile
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import Mock, patch
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import Client, TestCase, TransactionTestCase, override_settings
from django.utils import timezone
from core.models import Asset, CollectionItem, Edge, Job, KnowledgeBase, KnowledgeNode, NodeProposal
from core.services import auto_link_edges, link_collection_to_node, render_markdown, sync_node_to_file
from core.storage import bounded_path, save_upload
from core.document_parser import decode_text, parse
from core.ai_result import call
from core.jobs import enqueue, recover_expired
from office.models import WorkRecord
from bookshelf.models import Book


class IsolatedCase(TestCase):
    def setUp(self):
        self.workspace = tempfile.TemporaryDirectory()
        self.addCleanup(self.workspace.cleanup)
        root = Path(self.workspace.name)
        self.settings_override = override_settings(ZHI_SHI_ROOT=root / 'data', MEDIA_ROOT=root / 'media',
            BACKUP_ROOT=root / 'backup', ALLOWED_HOSTS=['testserver', 'localhost'])
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.base = KnowledgeBase.objects.create(name='测试库', directory='knowledge')
        self.user = get_user_model().objects.create_user('tester', password='test-password-123')
        self.admin = get_user_model().objects.create_user('administrator', password='test-password-123', is_staff=True)
        self.client.force_login(self.user)

    def node(self, **values):
        return KnowledgeNode.objects.create(base=self.base, title='测试', content_md='原始正文', **values)


class SecurityTests(IsolatedCase):
    def test_anonymous_business_pages_require_login(self):
        anonymous = Client()
        for url in ('/', '/collection/', '/settings/', '/knowledgebase/', '/bookshelf/', '/starmap/'):
            with self.subTest(url=url):
                self.assertEqual(anonymous.get(url).status_code, 302)

    def test_only_staff_can_read_settings_and_backup(self):
        self.assertEqual(self.client.get('/settings/').status_code, 403)
        self.assertEqual(self.client.get('/settings/backup/list/').status_code, 403)

    def test_staff_page_never_returns_token_and_never_fetches_gitee(self):
        from core.models import SystemConfig
        SystemConfig.set_value('gitee_token', 'fake-token-do-not-leak')
        self.client.force_login(self.admin)
        with patch('requests.get') as remote:
            response = self.client.get('/settings/')
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'fake-token-do-not-leak')
        remote.assert_not_called()

    def test_csrf_rejects_missing_token_and_wrong_origin(self):
        strict = Client(enforce_csrf_checks=True)
        strict.force_login(self.user)
        node = self.node()
        path = f'/knowledgebase/node/{node.pk}/note/'
        self.assertEqual(strict.post(path, {'note': '无令牌'}).status_code, 403)
        page = strict.get(f'/wiki/node/{node.pk}/')
        token = strict.cookies['csrftoken'].value
        self.assertEqual(strict.post(path, {'note': '正常', 'csrfmiddlewaretoken': token}).status_code, 302)
        self.assertEqual(strict.post(path, {'note': '错误来源'}, HTTP_X_CSRFTOKEN=token,
                                    HTTP_ORIGIN='https://untrusted.invalid').status_code, 403)

    def test_markdown_blocks_active_content_and_preserves_structure(self):
        output = render_markdown('<script>alert(1)</script><img src=x onerror=alert(2)>\n\n'
            '[bad](javascript:alert(3))\n\n```python\nprint("[[代码]]")\n```\n\n'
            '| 名称 | 值 |\n|---|---|\n|a|b|\n\n[[正常|显示]]')
        self.assertNotIn('<script', output)
        self.assertNotIn('onerror', output)
        self.assertNotIn('javascript:', output)
        self.assertIn('<table>', output)
        self.assertIn('<code', output)
        self.assertIn('[[代码]]', output)
        self.assertIn('kb-wikilink', output)

    def test_starmap_escapes_script_terminators(self):
        self.node().delete()
        KnowledgeNode.objects.create(base=self.base, title='</script><script>window.pwned=1</script>')
        response = self.client.get('/starmap/')
        self.assertNotContains(response, '</script><script>window.pwned')
        self.assertContains(response, 'application/json')
        self.assertEqual(response['X-Frame-Options'], 'SAMEORIGIN')

    def test_invalid_json_shapes_return_400(self):
        for body in ('[]', 'null', '{', '{"text":[]}'):
            with self.subTest(body=body):
                response = self.client.post('/knowledgebase/api/note/', body, content_type='application/json')
                self.assertEqual(response.status_code, 400)

    def test_cross_module_delete_is_rejected(self):
        node = self.node()
        self.assertEqual(self.client.post(f'/selfmedia/delete/{node.pk}/').status_code, 404)
        self.assertTrue(KnowledgeNode.objects.filter(pk=node.pk).exists())

    def test_unsafe_original_is_attachment_and_requires_login(self):
        from django.conf import settings
        Path(settings.MEDIA_ROOT).mkdir()
        (Path(settings.MEDIA_ROOT) / 'sample.html').write_text('<script>alert(1)</script>')
        self.assertEqual(Client().get('/media/sample.html').status_code, 302)
        response = self.client.get('/media/sample.html')
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('sandbox', response['Content-Security-Policy'])
        response.close()


class DataTests(IsolatedCase):
    def test_same_name_uploads_do_not_overwrite_and_identical_bytes_deduplicate(self):
        first, asset1 = save_upload(SimpleUploadedFile('same-name.txt', b'first'))
        second, asset2 = save_upload(SimpleUploadedFile('same-name.txt', b'second'))
        duplicate, asset3 = save_upload(SimpleUploadedFile('different-name.txt', b'first'))
        self.assertNotEqual(first, second)
        self.assertEqual(first.read_bytes(), b'first')
        self.assertEqual(second.read_bytes(), b'second')
        self.assertEqual(asset1.pk, asset3.pk)

    def test_invalid_upload_and_path_traversal_are_rejected(self):
        with self.assertRaises(ValueError):
            save_upload(SimpleUploadedFile('pretend.pdf', b'MZfake'))
        for relative in ('../adjacent/file', 'C:/outside/file', '/outside/file'):
            with self.subTest(relative=relative), self.assertRaises(ValueError):
                bounded_path(Path(self.workspace.name), relative)

    def test_export_identity_survives_rename_and_title_collision(self):
        first = self.node()
        second = self.node()
        path1, path2 = sync_node_to_file(first), sync_node_to_file(second)
        self.assertNotEqual(path1, path2)
        first.title = '新的名称'
        first.save()
        self.assertEqual(sync_node_to_file(first), path1)
        self.assertTrue(Path(path2).exists())

    def test_export_failure_is_recorded_and_import_is_idempotent(self):
        item = CollectionItem.objects.create(kind='text', title='幂等', parsed_md='正文', base=self.base)
        with patch('core.services.atomic_write', side_effect=OSError('disk full')):
            node = link_collection_to_node(item)
        self.assertEqual(node.export_status, 'failed')
        same = link_collection_to_node(item)
        self.assertEqual(same.pk, node.pk)
        self.assertEqual(KnowledgeNode.objects.count(), 1)
        self.assertEqual(same.export_status, 'done')

    def test_import_returns_failure_when_export_fails_but_preserves_original(self):
        with patch('core.services.atomic_write', side_effect=OSError('disk full')):
            response = self.client.post('/collection/import_local/',
                {'file': SimpleUploadedFile('same.txt', '中文原文'.encode())}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertFalse(response.json()['ok'])
        self.assertEqual(CollectionItem.objects.get().status, 'failed')
        self.assertEqual(Asset.objects.count(), 1)

    def test_relation_uniqueness_and_ambiguous_titles(self):
        source = self.node()
        target = KnowledgeNode.objects.create(base=self.base, title='目标')
        source.content_md = '[[目标]]'
        source.save()
        auto_link_edges(source)
        self.assertEqual(source.edges_out.count(), 1)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Edge.objects.create(source=source, target=target, kind='link')
        KnowledgeNode.objects.create(base=self.base, title='目标')
        auto_link_edges(source)
        self.assertEqual(source.edges_out.count(), 0)

    def test_archive_preserves_category_and_requires_completed_status(self):
        record = WorkRecord.objects.create(category='ops', content='网络维护', status='进行中')
        body = {'date': record.date.isoformat()}
        self.assertFalse(self.client.post('/office/api/archive_day/', body, content_type='application/json').json()['ok'])
        record.status = '已完成'
        record.save()
        self.assertTrue(self.client.post('/office/api/archive_day/', body, content_type='application/json').json()['ok'])
        record.refresh_from_db()
        self.assertEqual(record.category, 'ops')
        self.assertTrue(record.is_archived)

    def test_two_edit_windows_conflict_and_previous_content_is_recoverable(self):
        node = self.node()
        url = f'/wiki/node/{node.pk}/edit/'
        first = self.client.post(url, {'version': node.version, 'title': '修订', 'content_md': '新正文'})
        self.assertEqual(first.status_code, 302)
        second = self.client.post(url, {'version': node.version, 'title': '覆盖', 'content_md': '过期正文'})
        self.assertEqual(second.status_code, 409)
        node.refresh_from_db()
        self.assertEqual(node.content_md, '新正文')
        self.assertEqual(node.revisions.get().snapshot['content_md'], '原始正文')

    def test_ai_grow_requires_review_and_can_be_rejected(self):
        node = self.node()
        with patch('wiki.views.grow_node', return_value={'category': '候选分类', 'tags': '标签', 'definition': 'AI定义'}):
            response = self.client.post(f'/wiki/node/{node.pk}/grow/')
        node.refresh_from_db()
        self.assertEqual(node.category, '')
        proposal = NodeProposal.objects.get()
        self.client.post(response.json()['review_url'], {'action': 'reject'})
        node.refresh_from_db()
        self.assertEqual(node.content_md, '原始正文')

    def test_81st_node_is_reachable_by_pagination(self):
        KnowledgeNode.objects.bulk_create([KnowledgeNode(base=self.base, title=f'知识{i}') for i in range(81)])
        response = self.client.get('/knowledgebase/?page=2')
        self.assertEqual(response.context['page_obj'].paginator.count, 81)
        self.assertEqual(len(response.context['nodes']), 31)

    def test_default_search_never_calls_web_or_includes_pending(self):
        self.node(status='pending')
        with patch('office.views.search_web') as web:
            response = self.client.get('/office/search/?q=测试')
        web.assert_not_called()
        self.assertNotContains(response, '/wiki/node/')

    def test_statistics_query_count_is_independent_of_library_count(self):
        from django.test.utils import CaptureQueriesContext
        from django.db import connection
        with CaptureQueriesContext(connection) as initial:
            self.client.get('/knowledgebase/')
        KnowledgeBase.objects.bulk_create([KnowledgeBase(name=f'库{i}') for i in range(10)])
        with CaptureQueriesContext(connection) as larger:
            self.client.get('/knowledgebase/')
        self.assertLessEqual(len(larger), len(initial) + 2)


class ParserAndAITests(TestCase):
    def test_chinese_encoding_never_discards_bytes(self):
        for encoding in ('utf-8', 'gb18030', 'utf-16'):
            decoded, _ = decode_text('中文资料'.encode(encoding))
            self.assertEqual(decoded, '中文资料')
        with self.assertRaises(ValueError):
            decode_text(b'\xff\x81')

    def test_docx_paragraph_table_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph('表格前')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '列一'
        table.cell(1, 0).text = '表格内'
        document.add_paragraph('表格后')
        buffer = io.BytesIO()
        document.save(buffer)
        text, metadata = parse(buffer.getvalue(), 'sample.docx')
        self.assertLess(text.index('表格前'), text.index('表格内'))
        self.assertLess(text.index('表格内'), text.index('表格后'))
        self.assertTrue(metadata['warnings'])

    @patch.dict('os.environ', {'ZHISHIKU_ALLOW_AI_OUTBOUND':'1'})
    def test_ai_failure_types_and_no_automatic_forwarding(self):
        provider = {'id': 'test', 'api_key': 'fake', 'base_url': 'https://api.example.invalid/v1', 'model': 'test-model'}
        for status, expected in ((401, 'authentication_failed'), (429, 'rate_limited'), (503, 'upstream_error')):
            with self.subTest(status=status), patch('requests.post', return_value=Mock(status_code=status)):
                result = call(provider, 'system', 'input')
                self.assertFalse(result.ok)
                self.assertEqual(result.error_code, expected)
        malformed = Mock(status_code=200)
        malformed.json.return_value = {'choices': []}
        with patch('requests.post', return_value=malformed):
            self.assertEqual(call(provider, 'system', 'input').error_code, 'invalid_response')

    def test_local_ai_without_key_is_explicitly_supported(self):
        provider = {'id': 'local', 'local': True, 'auth_required': False,
                    'base_url': 'http://127.0.0.1:11434/v1', 'model': 'local-model'}
        response = Mock(status_code=200)
        response.json.return_value = {'choices': [{'message': {'content': '本地回答'}}]}
        with patch('requests.post', return_value=response) as transport:
            self.assertTrue(call(provider, 'system', 'input').ok)
            self.assertNotIn('Authorization', transport.call_args.kwargs['headers'])

    def test_ai_json_rejects_wrong_types(self):
        from core.extract import parse_ai_json
        self.assertIsNone(parse_ai_json('{"content_md":[],"title":5}'))
        self.assertIsNone(parse_ai_json('{"extra_links":[3]}', required_key=None))


class QueueTests(TestCase):
    def test_idempotency_and_restart_reconciliation(self):
        first = enqueue('video', {'item_id': 1}, key='sample')
        self.assertEqual(enqueue('video', {'item_id': 1}, key='sample').pk, first.pk)
        first.status = 'running'
        first.lease_until = timezone.now() - timedelta(seconds=1)
        first.save()
        recover_expired()
        first.refresh_from_db()
        self.assertEqual(first.status, 'failed')
        self.assertIsNone(first.active_key)
