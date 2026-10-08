"""Shared bounded parsing, with original-page evidence and explicit fidelity limits."""
from html import escape
from pathlib import Path
import io
from .storage import TEXT_EXTENSIONS

PARSER_VERSION = 'structured-text-2'


def decode_text(raw):
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        return raw.decode('utf-16'), 'utf-16'
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            pass
    raise ValueError('文本编码无法可靠识别；原件已保留，请转换为 UTF-8 后重试')


def table_markdown(rows):
    rows = [[str(c).replace('|', '\\|').replace('\n', '<br>') for c in row] for row in rows]
    if not rows:
        return ''
    width = max(map(len, rows))
    rows = [row + [''] * (width - len(row)) for row in rows]
    lines = ['| ' + ' | '.join(rows[0]) + ' |', '| ' + ' | '.join(['---'] * width) + ' |']
    lines.extend('| ' + ' | '.join(row) + ' |' for row in rows[1:])
    return '\n'.join(lines)


def parse(raw, name):
    ext = Path(name).suffix.lower()
    metadata = {'ext': ext, 'parser_version': PARSER_VERSION, 'warnings': []}
    if ext in TEXT_EXTENSIONS:
        text, encoding = decode_text(raw)
        metadata['encoding'] = encoding
        if ext in {'.md', '.markdown'}:
            return text, metadata
        return '<pre class="kb-plain-text">' + escape(text) + '</pre>', metadata
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('PDF 需要打开密码；原件已保留')
        if len(reader.pages) > 1000:
            raise ValueError('PDF 超过 1000 页解析上限；请使用原件阅读器')
        pages = [page.extract_text() or '' for page in reader.pages]
        if not any(text.strip() for text in pages):
            raise ValueError('PDF 无可提取文字层，可能是扫描件；原件已保留，需要 OCR')
        metadata.update(pages=len(pages), blocks=[{'page': i, 'type': 'preformatted_text', 'text': text}
                                                  for i, text in enumerate(pages, 1)])
        metadata['warnings'].append('PDF 为逐页文字提取；多栏顺序、图片、表格和公式请对照原件，未声称完整重排保真')
        return '\n\n'.join(f'## 原件第 {i} 页\n\n<pre class="kb-plain-text">{escape(text)}</pre>'
                            for i, text in enumerate(pages, 1)), metadata
    if ext == '.docx':
        from docx import Document
        from docx.oxml.text.paragraph import CT_P
        from docx.oxml.table import CT_Tbl
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        document = Document(io.BytesIO(raw))
        blocks = []
        for element in document.element.body.iterchildren():
            if isinstance(element, CT_P):
                paragraph = Paragraph(element, document)
                text = paragraph.text
                if paragraph.style and paragraph.style.name.startswith('Heading '):
                    try:
                        text = '#' * min(6, int(paragraph.style.name.split()[-1])) + ' ' + text
                    except ValueError:
                        pass
                if text.strip():
                    blocks.append(text)
            elif isinstance(element, CT_Tbl):
                table = Table(element, document)
                blocks.append(table_markdown([[cell.text for cell in row.cells] for row in table.rows]))
        metadata['warnings'].append('DOCX 保留段落/表格顺序；嵌入图片、公式和复杂样式请查看原件')
        return '\n\n'.join(blocks), metadata
    if ext == '.xlsx':
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            blocks = []
            count = 0
            for sheet in workbook:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    count += len(row)
                    if count > 100000:
                        raise ValueError('表格超过十万单元格解析上限')
                    rows.append(['' if cell is None else str(cell) for cell in row])
                blocks.extend(['## ' + sheet.title, table_markdown(rows)])
            return '\n\n'.join(blocks), metadata
        finally:
            workbook.close()
    if ext == '.pptx':
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(raw))
        if len(presentation.slides) > 1000:
            raise ValueError('PPT 超过 1000 页解析上限')
        blocks = []
        for number, slide in enumerate(presentation.slides, 1):
            blocks.append('## 原件第 %d 页' % number)
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    blocks.append(shape.text_frame.text.strip())
                if shape.has_table:
                    blocks.append(table_markdown([[cell.text for cell in row.cells] for row in shape.table.rows]))
        metadata.update(slides=len(presentation.slides))
        metadata['warnings'].append('PPT 为文字/表格提取；图表、图片、公式和复杂阅读顺序请对照原件')
        if len(blocks) == len(presentation.slides):
            raise ValueError('PPT 未提取到正文，原件已保留')
        return '\n\n'.join(blocks), metadata
    return None, metadata
