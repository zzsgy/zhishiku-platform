"""书架：收集书籍、分类展示、阅读进度、读书笔记、划句问 AI、金句库。

书籍元信息同步写入知识库 03_知识库/06_书籍；正文支持 txt/md/各类文本/PDF/DOCX/XLSX/HTML 及网页链接，可在阅读器内直接阅读并划句。
"""
import json
import os
import logging
from django.shortcuts import render, redirect, get_object_or_404

logger = logging.getLogger('kb')
from django.conf import settings
from django.contrib import messages
from django.http import JsonResponse
from django.db.models import Q
from django.views.decorators.csrf import csrf_exempt

from .models import Book, ReadingNote, GoldenSentence
from core.services import log_operation, render_markdown, zhi_shi_path, ensure_dir, sanitize_filename
from core.ai import ask_ai, ai_available


def _export_book(book):
    """把书籍元数据作为 Markdown 落到知识库 06_书籍，便于统一检索。"""
    d = zhi_shi_path('03_知识库', '06_书籍')
    ensure_dir(d)
    content = (
        f'# {book.title}\n\n'
        f'- 作者：{book.author or "—"}\n'
        f'- 分类：{book.category or "—"}\n'
        f'- 总页数：{book.total_pages}\n'
        f'- 当前进度：{book.progress}\n'
        f'- 状态：{book.get_status_display()}\n'
    )
    try:
        (d / f'{sanitize_filename(book.title)}.md').write_text(content, encoding='utf-8')
    except Exception:
        pass


def index(request):
    cat = (request.GET.get('cat') or '').strip()
    q = (request.GET.get('q') or '').strip()
    err = request.GET.get('err')
    books = Book.objects.all()
    if cat:
        books = books.filter(category=cat)
    if q:
        books = books.filter(title__icontains=q)
    categories = [c for c in Book.objects.values_list('category', flat=True).distinct() if c]
    return render(request, 'bookshelf.html', {
        'books': books, 'categories': categories, 'cat': cat, 'q': q, 'err': err,
    })


def add(request):
    if request.method != 'POST':
        return redirect('/bookshelf/')
    f = request.FILES.get('file')
    url = (request.POST.get('url') or '').strip()
    # 前置校验：链接与文件至少其一，避免误触产生空白书籍（与「知识收集」导入行为一致）
    if not f and not url:
        return redirect('/bookshelf/?err=1')
    title = (request.POST.get('title') or '').strip()
    author = (request.POST.get('author') or '').strip()
    category = (request.POST.get('category') or '').strip()
    try:
        total = int(request.POST.get('total_pages') or 0)
    except ValueError:
        total = 0
    file_path = ''
    source_url = ''
    if f:
        # 本地文件：支持各类文本 / PDF / DOCX / XLSX / 网页 等
        d = settings.MEDIA_ROOT / 'books'
        d.mkdir(parents=True, exist_ok=True)
        dest = d / f.name
        with open(dest, 'wb') as out:
            for chunk in f.chunks():
                out.write(chunk)
        file_path = str(dest)
    elif url:
        # 网页链接：抓取正文转 Markdown 落盘，失败则归档至链接库
        source_url = url
        try:
            from core.parsers import parse_web_page
            t, md = parse_web_page(url, 'auto')
            d = settings.MEDIA_ROOT / 'books'
            d.mkdir(parents=True, exist_ok=True)
            safe = sanitize_filename(title or t)[:80]
            md_path = d / f'{safe}.md'
            md_path.write_text(md, encoding='utf-8')
            file_path = str(md_path)
            title = title or t
        except Exception as e:
            try:
                from collection.views import record_failed_link
                record_failed_link(url, 'web', 'unknown', link_type='web', title=title or url)
            except Exception:
                pass
            messages.warning(request, f'网页抓取失败（已归档至链接库）：{e}')
            return redirect('/bookshelf/?err=web')
    book = Book.objects.create(
        title=title or (f.name if f else (url if url else '未命名书籍')),
        author=author, category=category, total_pages=total,
        file_path=file_path, source_url=source_url)
    _export_book(book)
    log_operation('bookshelf', 'add', detail=book.title)
    return redirect('/bookshelf/')


# 可作为纯文本直接阅读的扩展名（其余交给专用解析器）
TEXT_EXTS = ('.txt', '.text', '.csv', '.tsv', '.json', '.log', '.xml', '.yaml',
             '.yml', '.py', '.js', '.css', '.sql', '.ini', '.cfg', '.rtf',
             '.html', '.htm')


def _read_book_body(book):
    """根据文件类型读取书籍正文，返回渲染后的 HTML；不支持或读取失败时返回 ''。"""
    fp = book.file_path
    if not fp or not os.path.exists(fp):
        return ''
    p = fp.lower()
    try:
        if p.endswith(('.md', '.markdown')):
            raw = open(fp, encoding='utf-8', errors='ignore').read()
            return render_markdown(raw)
        if p.endswith(TEXT_EXTS):
            raw = open(fp, encoding='utf-8', errors='ignore').read()
            if p.endswith(('.html', '.htm')):
                from bs4 import BeautifulSoup
                soup = BeautifulSoup(raw, 'html.parser')
                for tag in soup(['script', 'style']):
                    tag.decompose()
                raw = soup.get_text('\n', strip=True)
            return f'<pre class="kb-prose">{raw}</pre>'
        if p.endswith('.pdf'):
            import PyPDF2
            reader = PyPDF2.PdfReader(fp)
            if reader.is_encrypted:
                try:
                    reader.decrypt('')  # 空密码解密（仅限作者未设打开密码的 PDF）
                except Exception:
                    return ('<p class="text-muted">该 PDF 受密码保护，无法在线提取正文，'
                            '请用本地阅读器打开。</p>')
            text = '\n\n'.join((pg.extract_text() or '') for pg in reader.pages)
            if not text.strip():
                return ('<p class="text-muted">该 PDF 未提取到文本：可能是扫描件/图片型 PDF'
                        '（无文字层），暂不支持 OCR，请用本地阅读器打开。</p>')
            return f'<pre class="kb-prose">{text}</pre>'
        if p.endswith('.docx'):
            from docx import Document
            doc = Document(fp)
            lines = [para.text for para in doc.paragraphs if para.text.strip()]
            content = '\n'.join(lines)
            return f'<pre class="kb-prose">{content}</pre>'
        if p.endswith('.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f'# {ws.title}')
                for row in ws.iter_rows(values_only=True):
                    out.append('\t'.join('' if c is None else str(c) for c in row))
            content = '\n'.join(out)
            return f'<pre class="kb-prose">{content}</pre>'
    except Exception as e:
        logger.error('read book body failed: %s', e)
    return ''


def book_view(request, pk):
    book = get_object_or_404(Book, pk=pk)
    notes = book.notes.all()
    # 本书金句：优先显式归属(book FK)，兼容旧数据按 source 标题匹配
    goldens = GoldenSentence.objects.filter(
        Q(book=book) | Q(book__isnull=True, source__icontains=book.title)
    ).distinct()
    body = _read_book_body(book)
    return render(request, 'book_reader.html', {
        'book': book, 'notes': notes, 'goldens': goldens, 'body': body,
    })


def book_retry(request, pk):
    """对来源为网页链接的书籍重新抓取并更新正文。"""
    if request.method != 'POST':
        return redirect(f'/bookshelf/book/{pk}/')
    book = get_object_or_404(Book, pk=pk)
    if not book.source_url:
        return redirect(f'/bookshelf/book/{pk}/')
    try:
        from core.parsers import parse_web_page
        t, md = parse_web_page(book.source_url, 'auto')
        d = settings.MEDIA_ROOT / 'books'
        d.mkdir(parents=True, exist_ok=True)
        safe = sanitize_filename(book.title)[:80]
        md_path = d / f'{safe}.md'
        md_path.write_text(md, encoding='utf-8')
        book.file_path = str(md_path)
        book.save(update_fields=['file_path'])
        messages.success(request, f'已重新抓取并更新正文：{t}')
        log_operation('bookshelf', 'retry', detail=book.title)
    except Exception as e:
        messages.error(request, f'重新抓取失败：{e}')
    return redirect(f'/bookshelf/book/{pk}/')


def update_progress(request, pk):
    if request.method != 'POST':
        return redirect('/bookshelf/')
    book = get_object_or_404(Book, pk=pk)
    p = int(request.POST.get('progress') or 0)
    book.progress = p
    if book.total_pages and p >= book.total_pages:
        book.status = 'done'
    elif p > 0:
        book.status = 'reading'
    else:
        book.status = 'unread'
    book.save()
    _export_book(book)
    log_operation('bookshelf', 'progress', detail=f'{book.title} -> {p}')
    return redirect(f'/bookshelf/book/{pk}/')


def add_note(request, pk):
    if request.method != 'POST':
        return redirect('/bookshelf/')
    book = get_object_or_404(Book, pk=pk)
    loc = (request.POST.get('location') or '').strip()
    note = (request.POST.get('note') or '').strip()
    if note:
        ReadingNote.objects.create(book=book, location=loc, note=note)
        log_operation('bookshelf', 'note', detail=book.title)
    return redirect(f'/bookshelf/book/{pk}/')


@csrf_exempt
def api_golden(request):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'POST required'})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'ok': False, 'error': 'invalid json'})
    text = (data.get('text') or '').strip()
    source = (data.get('source') or '').strip()
    if not text:
        return JsonResponse({'ok': False, 'error': '内容为空'})
    obj = GoldenSentence(text=text, source=source)
    book_id = data.get('book')
    node_id = data.get('node')
    if book_id:
        try:
            obj.book_id = int(book_id)
        except (TypeError, ValueError):
            pass
    if node_id:
        try:
            obj.node_id = int(node_id)
        except (TypeError, ValueError):
            pass
    obj.save()
    log_operation('bookshelf', 'golden', detail=source)
    return JsonResponse({'ok': True, 'id': obj.pk})


@csrf_exempt
def api_golden_delete(request):
    """删除金句（阅读面板手动移除，仅此途径才移除已收录金句）。"""
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'POST required'})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'ok': False, 'error': 'invalid json'})
    gid = data.get('id')
    if not gid:
        return JsonResponse({'ok': False, 'error': '缺少 id'})
    GoldenSentence.objects.filter(pk=gid).delete()
    log_operation('bookshelf', 'golden_delete', detail='id=%s' % gid)
    return JsonResponse({'ok': True})


@csrf_exempt
def api_note_delete(request):
    """删除读书笔记（阅读面板手动移除）。"""
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'POST required'})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'ok': False, 'error': 'invalid json'})
    nid = data.get('id')
    if not nid:
        return JsonResponse({'ok': False, 'error': '缺少 id'})
    ReadingNote.objects.filter(pk=nid).delete()
    log_operation('bookshelf', 'note_delete', detail='id=%s' % nid)
    return JsonResponse({'ok': True})


def golden_list(request):
    items = GoldenSentence.objects.all()[:200]
    return render(request, 'golden.html', {'items': items})


def api_note(request):
    """划词「记笔记」JSON 接口：把选中内容快速写入读书笔记。"""
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'POST required'})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'ok': False, 'error': 'invalid json'})
    book = Book.objects.filter(pk=data.get('book')).first()
    if not book:
        return JsonResponse({'ok': False, 'error': 'book not found'})
    note = (data.get('note') or '').strip()
    if not note:
        return JsonResponse({'ok': False, 'error': '内容为空'})
    obj = ReadingNote.objects.create(
        book=book, location=(data.get('location') or '').strip(), note=note)
    log_operation('bookshelf', 'note', detail=book.title)
    return JsonResponse({'ok': True, 'id': obj.pk})


@csrf_exempt
def api_ask(request):
    """书籍阅读页划词问 AI：携带书名与选中文本作为上下文，调用现有 AI 模块。"""
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'POST required'})
    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'ok': False, 'error': 'invalid json'})
    text = (data.get('text') or '').strip()
    question = (data.get('question') or '').strip()
    provider = (data.get('provider') or '').strip() or None
    if not text or not question:
        return JsonResponse({'ok': False, 'error': '缺少文本或问题'})
    book = Book.objects.filter(pk=data.get('book')).first()
    book_title = book.title if book else ''
    system = ('你是个人知识库助手，请基于用户提供的资料（来自书籍）作答，要求：'
              '简洁、准确、用中文，必要时给出可执行的建议。')
    user_prompt = (
        f"资料（来自书籍《{book_title}》）：\n{text}\n\n问题：{question}"
        if book_title else f'资料：\n{text}\n\n问题：{question}'
    )
    ans = ask_ai(system, user_prompt, provider_id=provider)
    if not ans:
        ans = ('（当前未配置可用的 AI 服务，无法调用 AI。请在「系统设置 → AI 服务配置』'
               '填写对应服务的 API Key 并启用。）') if not ai_available() else '（调用失败，请检查该服务的 Key / 网络）'
    return JsonResponse({'ok': True, 'answer': ans})


def book_delete(request, pk):
    """删除书籍：级联删除读书笔记，并安全删除其本地文件（仅限 MEDIA_ROOT 内）。"""
    if request.method != 'POST':
        return redirect('/bookshelf/')
    book = get_object_or_404(Book, pk=pk)
    title = book.title
    fp = book.file_path
    if fp:
        try:
            abspath = os.path.abspath(fp)
            root = os.path.abspath(str(settings.MEDIA_ROOT))
            if abspath.startswith(root) and os.path.exists(abspath):
                os.remove(abspath)
        except OSError:
            pass
    book.delete()  # 级联删除 ReadingNote
    log_operation('bookshelf', 'delete', detail=title)
    return redirect('/bookshelf/')
