import io
import json
from pathlib import Path
from django.test import Client
from docx import Document
from PIL import Image
from pypdf import PdfWriter
from core.tests.test_regressions import IsolatedCase
from bookshelf.models import Book, ReaderPosition
from bookshelf.reader import preview
from bookshelf.reader_views import file_key


class ReaderTests(IsolatedCase):
    def book(self, extension, raw):
        path = self.workspace.name + '/media/sample' + extension
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        Path(path).write_bytes(raw)
        return Book.objects.create(title='阅读样本', file_path='sample' + extension)

    def test_original_pdf_inline_ranges_auth_and_immutable_download(self):
        stream = io.BytesIO()
        writer = PdfWriter(); writer.add_blank_page(width=595,height=842); writer.write(stream)
        raw = stream.getvalue(); book = self.book('.pdf',raw)
        url = f'/bookshelf/book/{book.pk}/original/'
        self.assertEqual(Client().get(url).status_code,302)
        response = self.client.get(url)
        self.assertEqual(response['Content-Type'],'application/pdf')
        self.assertTrue(response['Content-Disposition'].startswith('inline'))
        self.assertEqual(b''.join(response.streaming_content),raw)
        response = self.client.get(url,HTTP_RANGE='bytes=0-9')
        self.assertEqual(response.status_code,206)
        self.assertEqual(b''.join(response.streaming_content),raw[:10])
        self.assertEqual(response['Content-Range'],f'bytes 0-9/{len(raw)}')
        response = self.client.get(url,HTTP_RANGE='bytes=-7')
        self.assertEqual(b''.join(response.streaming_content),raw[-7:])
        self.assertEqual(self.client.get(url,HTTP_RANGE='bytes=999999-').status_code,416)
        self.assertEqual(self.client.get(url,HTTP_RANGE='bytes=0-2,8-9').status_code,416)
        self.assertEqual(self.client.head(url)['Content-Length'],str(len(raw)))
        response = self.client.get(url+'?download=1')
        self.assertTrue(response['Content-Disposition'].startswith('attachment'))
        response.close()
        self.assertEqual(Path(self.workspace.name+'/media/sample.pdf').read_bytes(),raw)

    def test_reader_never_flattens_pdf_or_requires_extracted_text(self):
        stream=io.BytesIO();writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.write(stream)
        book=self.book('.pdf',stream.getvalue())
        response=self.client.get(f'/bookshelf/book/{book.pk}/')
        self.assertContains(response,'id="pdfContainer"')
        self.assertNotContains(response,'扫描件需 OCR 后检索')
        self.assertEqual(ReaderPosition.objects.count(),0)

    def test_word_preserves_heading_run_image_table_and_order(self):
        doc=Document();doc.add_heading('标题',1);doc.add_paragraph('图前正文')
        doc.add_paragraph('第一项',style='List Number');doc.add_paragraph('第二项',style='List Number')
        paragraph=doc.add_paragraph();paragraph.add_run('<script>不能执行</script>').bold=True
        image=io.BytesIO();Image.new('RGB',(20,20),'white').save(image,format='PNG');image.seek(0)
        doc.add_picture(image);doc.add_paragraph('图后正文')
        table=doc.add_table(rows=2,cols=2);table.cell(0,0).merge(table.cell(0,1)).text='合并标题'
        table.cell(1,0).text='左格';table.cell(1,1).text='右格';doc.add_paragraph('表后正文')
        stream=io.BytesIO();doc.save(stream);book=self.book('.docx',stream.getvalue())
        body=preview(book)['html']
        self.assertIn('<h1>标题</h1>',body);self.assertIn('<strong>&lt;script&gt;',body)
        self.assertIn('colspan="2"',body);self.assertNotIn('<script>',body)
        self.assertIn('reader-list-marker">1.',body);self.assertIn('reader-list-marker">2.',body)
        self.assertLess(body.index('图前正文'),body.index('<img'));self.assertLess(body.index('<img'),body.index('图后正文'))
        self.assertLess(body.index('图后正文'),body.index('<table>'));self.assertLess(body.index('</table>'),body.index('表后正文'))
        import re
        image_url=re.search('src="([^"]+)"',body).group(1)
        self.assertEqual(self.client.get(image_url)['Content-Type'],'image/png')
        self.assertEqual(Client().get(image_url).status_code,302)

    def test_html_scripts_events_and_unsafe_links_removed(self):
        book=self.book('.html',b'<h1>Title</h1><script>alert(1)</script><img src=x onerror=alert(2)><a href="javascript:alert(3)">x</a><table><tr><td>cell</td></tr></table>')
        html=preview(book)['html']
        for unsafe in ['<script','onerror','javascript:','alert(1)']:self.assertNotIn(unsafe,html)
        self.assertIn('<table>',html);self.assertIn('<h1>Title',html)

    def test_text_encodings_and_code_indentation(self):
        book=self.book('.txt','第一行\n第二行\n\n新段落'.encode('utf-16'))
        html=preview(book)['html'];self.assertIn('第一行\n第二行',html);self.assertEqual(html.count('<p'),2)
        book=self.book('.py',b'if True:\n    print("<script>")\n')
        html=preview(book)['html'];self.assertIn('    print',html);self.assertIn('<pre><code>',html);self.assertNotIn('<script>',html)

    def test_csv_preserves_quoted_cells_and_chinese(self):
        book=self.book('.csv','名称,说明\n中文,"一,二"\n'.encode('gb18030'))
        html=preview(book)['html'];self.assertIn('<td>中文</td>',html);self.assertIn('<td>一,二</td>',html)

    def test_rtf_unicode_and_old_doc_fallback_keep_original(self):
        book=self.book('.rtf',br'{\rtf1\ansi First\par\u20013?\u25991?}')
        html=preview(book)['html'];self.assertIn('中文',html);self.assertNotIn('\\rtf1',html)
        raw=b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'+b'\x00'*64
        book=self.book('.doc',raw)
        response=self.client.get(f'/bookshelf/book/{book.pk}/preview/')
        self.assertEqual(response.status_code,422);self.assertIn('另存为 DOCX',response.json()['error'])
        response=self.client.get(f'/bookshelf/book/{book.pk}/original/?download=1')
        self.assertEqual(b''.join(response.streaming_content),raw)

    def test_position_per_user_without_rewriting_book_or_original(self):
        book=self.book('.txt',b'original');before=Book.objects.values().get(pk=book.pk)
        data={'fileKey':file_key(book),'page':1,'offset':.5,'percent':50,'fontSize':20,'paper':'warm'}
        url=f'/bookshelf/book/{book.pk}/position/'
        self.assertEqual(self.client.post(url,json.dumps(data),content_type='application/json').status_code,200)
        self.client.force_login(self.admin)
        response=self.client.get(f'/bookshelf/book/{book.pk}/')
        self.assertFalse(response.context['reader_config']['hasPosition'])
        self.client.force_login(self.user)
        response=self.client.get(f'/bookshelf/book/{book.pk}/');self.assertEqual(response.context['reader_config']['offset'],.5)
        self.assertEqual(Book.objects.values().get(pk=book.pk),before)
        self.assertEqual(Path(self.workspace.name+'/media/sample.txt').read_bytes(),b'original')
        Path(self.workspace.name+'/media/sample.txt').write_bytes(b'changed original')
        self.assertEqual(self.client.post(url,json.dumps(data),content_type='application/json').status_code,409)
        self.assertFalse(self.client.get(f'/bookshelf/book/{book.pk}/').context['reader_config']['hasPosition'])

    def test_position_rejects_nan_ranges_and_missing_csrf(self):
        book=self.book('.txt',b'original');url=f'/bookshelf/book/{book.pk}/position/'
        valid={'fileKey':file_key(book),'page':1,'offset':0,'percent':0}
        for values in [{'offset':float('nan')},{'percent':101},{'page':0},{'fontSize':100},{'page':True}]:
            self.assertEqual(self.client.post(url,json.dumps({**valid,**values}),content_type='application/json').status_code,400)
        strict=Client(enforce_csrf_checks=True);strict.force_login(self.user)
        self.assertEqual(strict.post(url,json.dumps(valid),content_type='application/json').status_code,403)

    def test_legacy_paths_are_bounded_and_missing_original_has_fallback(self):
        book=Book.objects.create(title='错误路径',file_path='../outside.pdf')
        self.assertEqual(self.client.get(f'/bookshelf/book/{book.pk}/original/').status_code,404)
        self.assertFalse(self.client.get(f'/bookshelf/book/{book.pk}/').context['reader_config']['available'])
        self.assertEqual(self.client.get(f'/bookshelf/book/{book.pk}/preview/').status_code,422)
